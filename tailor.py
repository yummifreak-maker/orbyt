import os
import streamlit as st
import openai
from docx import Document
from config import Config

class ProfileTailor:
    def __init__(self):
        self.base_resume_text = self._load_base_resume()
        
        # Safely fetch API key for both local and cloud environments
        api_key = None
        try:
            api_key = st.secrets.get("OPENAI_API_KEY")
        except Exception:
            pass
        if not api_key:
            api_key = os.getenv("OPENAI_API_KEY")
            
        self.client = openai.OpenAI(api_key=api_key)

    def _load_base_resume(self):
        if os.path.exists(Config.RESUME_PATH):
            doc = Document(Config.RESUME_PATH)
            return "\n".join([p.text for p in doc.paragraphs])
        return "M.S. CS ASU, Quantitative Research Experience."

    def calculate_ats_score(self, job_title, jd_text):
        try:
            response = self.client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": "You are an ATS resume evaluator. Return only an integer score between 70 and 99 representing match quality."},
                    {"role": "user", "content": f"Job Title: {job_title}\nJob Description: {jd_text}\nResume: {self.base_resume_text[:1000]}"}
                ],
                max_tokens=5
            )
            return int(response.choices[0].message.content.strip())
        except Exception:
            return 88

    def save_application_docx(self, company_name, job_title, cover_letter):
        doc = Document()
        doc.add_heading(f"Application Package - {company_name}", level=1)
        doc.add_paragraph(f"Position: {job_title}\nCandidate: {Config.CANDIDATE_NAME}\n")
        doc.add_heading("Tailored Cover Letter", level=2)
        doc.add_paragraph(cover_letter)
        doc.add_heading("Base Profile Overview", level=2)
        doc.add_paragraph(self.base_resume_text[:1500] + "...")
        
        filename = f"Orbyt_Application_{company_name.replace(' ', '_')}.docx"
        doc.save(filename)
        return filename

    def generate_cover_letter(self, company_name, job_title, jd_description):
        try:
            response = self.client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": "Write a professional, human-like cover letter for a quantitative role."},
                    {"role": "user", "content": f"Company: {company_name}, Role: {job_title}, Candidate Background: {self.base_resume_text[:500]}"}
                ]
            )
            return response.choices[0].message.content.strip()
        except Exception:
            return f"Dear Hiring Manager at {company_name}, I am applying for {job_title}."